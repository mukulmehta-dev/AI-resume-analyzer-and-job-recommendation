import io
import os
import zipfile
import fitz
import docx
from fastapi import HTTPException, status
from backend.app.core.config import settings


def validate_file_content(content: bytes, filename: str) -> str:
    max_bytes = settings.MAX_UPLOAD_MB * 1024 * 1024
    if len(content) > max_bytes:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail=f"File size exceeds maximum allowed limit of {settings.MAX_UPLOAD_MB}MB"
        )

    _, ext = os.path.splitext(filename)
    ext = ext.lower()

    if ext not in [".pdf", ".docx"]:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Unsupported file format. Only PDF (.pdf) and DOCX (.docx) documents are permitted."
        )

    if ext == ".pdf":
        if not content.startswith(b"%PDF"):
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail="Invalid file signature: file does not match genuine PDF format."
            )
        return "application/pdf"

    if ext == ".docx":
        if not content.startswith(b"PK\x03\x04"):
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail="Invalid file signature: file does not match genuine DOCX format."
            )
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                if "word/document.xml" not in zf.namelist():
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                        detail="Corrupted DOCX structure: missing word/document.xml entry."
                    )
                total_uncompressed = sum(info.file_size for info in zf.infolist())
                if total_uncompressed > 50 * 1024 * 1024:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                        detail="Uncompressed document size exceeds maximum allowed limit of 50MB."
                    )
                if len(content) > 0 and (total_uncompressed / len(content)) > 100:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                        detail="Suspicious compression ratio detected in DOCX archive."
                    )
        except zipfile.BadZipFile:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                detail="Corrupt DOCX archive: unable to read document structure."
            )
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    raise HTTPException(
        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
        detail="Unsupported file format."
    )


def extract_text_from_pdf(content: bytes) -> str:
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Corrupt PDF file: unable to parse document stream."
        )

    if doc.page_count == 0:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Empty PDF document: no pages found."
        )

    if doc.page_count > 30:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="PDF document exceeds maximum allowed length of 30 pages."
        )

    text_parts = []
    for page in doc:
        text_parts.append(page.get_text())

    full_text = "\n".join(text_parts).strip()
    if not full_text or len(full_text) < 20:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Scanned-only or empty document: selectable text could not be extracted from PDF."
        )
    return full_text


def extract_text_from_docx(content: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(content))
    except Exception:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Corrupt DOCX file: unable to parse document body."
        )

    text_parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text:
            text_parts.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                text_parts.append(" | ".join(row_texts))

    full_text = "\n".join(text_parts).strip()
    if not full_text or len(full_text) < 20:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail="Empty DOCX document: no text content detected."
        )
    return full_text


def extract_resume_text(content: bytes, filename: str) -> str:
    _, ext = os.path.splitext(filename)
    ext = ext.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(content)
    if ext == ".docx":
        return extract_text_from_docx(content)
    raise HTTPException(
        status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
        detail="Unsupported file format."
    )
