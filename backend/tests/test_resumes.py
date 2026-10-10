import os
import io
import zipfile
import fitz
import docx
from backend.app.models.resume import Resume


def create_sample_pdf(page_count: int = 1) -> bytes:
    doc = fitz.open()
    for _ in range(page_count):
        page = doc.new_page()
        page.insert_text((50, 72), "Alex Morgan\nSoftware Engineer\nExpert in Python, FastAPI, React, SQL, and Docker development.")
    return doc.tobytes()


def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Jordan Taylor - Data Scientist", level=1)
    doc.add_paragraph("Skilled in Python, PyTorch, Pandas, Scikit-learn, SQL, and Machine Learning workflows.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_upload_valid_pdf_and_list(client):
    reg = client.post("/api/auth/register", json={
        "name": "PDF User",
        "email": "pdfuser@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    pdf_bytes = create_sample_pdf()
    files = {"file": ("resume.pdf", pdf_bytes, "application/pdf")}
    res = client.post("/api/resumes", files=files)
    assert res.status_code == 201
    data = res.json()
    assert data["filename"] == "resume.pdf"
    assert data["mime"] == "application/pdf"
    assert "Alex Morgan" in data["extracted_text"]
    assert "FastAPI" in data["extracted_text"]

    resume_id = data["id"]
    list_res = client.get("/api/resumes")
    assert list_res.status_code == 200
    resumes_list = list_res.json()
    assert len(resumes_list) == 1
    assert resumes_list[0]["id"] == resume_id


def test_upload_valid_docx(client):
    reg = client.post("/api/auth/register", json={
        "name": "Docx User",
        "email": "docxuser@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    docx_bytes = create_sample_docx()
    files = {"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/resumes", files=files)
    assert res.status_code == 201
    data = res.json()
    assert data["filename"] == "resume.docx"
    assert "Jordan Taylor" in data["extracted_text"]
    assert "Machine Learning" in data["extracted_text"]


def test_exe_renamed_to_pdf_rejected(client):
    reg = client.post("/api/auth/register", json={
        "name": "Exe User",
        "email": "exeuser@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    fake_exe = b"MZ\x90\x00\x03\x00\x00\x00" + b"This is a Windows executable binary payload."
    files = {"file": ("malicious.pdf", fake_exe, "application/pdf")}
    res = client.post("/api/resumes", files=files)
    assert res.status_code == 422
    assert "invalid file signature" in res.json()["detail"].lower()


def test_oversize_file_rejected(client):
    reg = client.post("/api/auth/register", json={
        "name": "Large User",
        "email": "largeuser@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    oversize_bytes = b"%PDF-1.4 " + (b"0" * (6 * 1024 * 1024))
    files = {"file": ("large.pdf", oversize_bytes, "application/pdf")}
    res = client.post("/api/resumes", files=files)
    assert res.status_code == 422
    assert "exceeds maximum allowed limit" in res.json()["detail"].lower()


def test_corrupt_pdf_rejected(client):
    reg = client.post("/api/auth/register", json={
        "name": "Corrupt User",
        "email": "corruptuser@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    corrupt_bytes = b"%PDF-1.4\ncorrupted header and corrupted body that cannot be parsed by fitz"
    files = {"file": ("corrupt.pdf", corrupt_bytes, "application/pdf")}
    res = client.post("/api/resumes", files=files)
    assert res.status_code == 422
    assert "corrupt" in res.json()["detail"].lower() or "scanned" in res.json()["detail"].lower()


def test_pdf_exceeding_30_pages_rejected(client):
    reg = client.post("/api/auth/register", json={
        "name": "Long PDF User",
        "email": "longpdf@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    long_pdf = create_sample_pdf(page_count=31)
    files = {"file": ("long.pdf", long_pdf, "application/pdf")}
    res = client.post("/api/resumes", files=files)
    assert res.status_code == 422
    assert "exceeds maximum allowed length of 30 pages" in res.json()["detail"]


def test_cross_user_access_returns_404(client):
    reg1 = client.post("/api/auth/register", json={
        "name": "User One",
        "email": "owner@example.com",
        "password": "password123"
    })
    assert reg1.status_code == 201
    pdf_bytes = create_sample_pdf()
    files = {"file": ("owner_resume.pdf", pdf_bytes, "application/pdf")}
    upload_res = client.post("/api/resumes", files=files)
    resume_id = upload_res.json()["id"]

    client.post("/api/auth/logout")

    reg2 = client.post("/api/auth/register", json={
        "name": "User Two",
        "email": "stranger@example.com",
        "password": "password123"
    })
    assert reg2.status_code == 201

    get_res = client.get(f"/api/resumes/{resume_id}")
    assert get_res.status_code == 404

    stream_res = client.get(f"/api/resumes/{resume_id}/file")
    assert stream_res.status_code == 404

    del_res = client.delete(f"/api/resumes/{resume_id}")
    assert del_res.status_code == 404


def test_delete_removes_file_from_disk(client, db_session):
    reg = client.post("/api/auth/register", json={
        "name": "Delete Tester",
        "email": "deleteuser@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    pdf_bytes = create_sample_pdf()
    files = {"file": ("delete_test.pdf", pdf_bytes, "application/pdf")}
    upload_res = client.post("/api/resumes", files=files)
    assert upload_res.status_code == 201
    resume_id = upload_res.json()["id"]

    resume_record = db_session.query(Resume).filter(Resume.id == resume_id).first()
    assert resume_record is not None
    stored_path = resume_record.stored_path
    assert os.path.exists(stored_path)

    del_res = client.delete(f"/api/resumes/{resume_id}")
    assert del_res.status_code == 200
    assert not os.path.exists(stored_path)

    fetch_deleted = client.get(f"/api/resumes/{resume_id}")
    assert fetch_deleted.status_code == 404


def test_file_streaming_security_headers(client):
    reg = client.post("/api/auth/register", json={
        "name": "Security Tester",
        "email": "sectester@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    pdf_bytes = create_sample_pdf()
    files = {"file": ("my resume (v1).pdf", pdf_bytes, "application/pdf")}
    upload_res = client.post("/api/resumes", files=files)
    assert upload_res.status_code == 201
    resume_id = upload_res.json()["id"]

    file_res = client.get(f"/api/resumes/{resume_id}/file")
    assert file_res.status_code == 200
    assert file_res.headers.get("x-content-type-options") == "nosniff"
    assert file_res.headers.get("content-security-policy") == "sandbox"
    assert 'filename="my_resume__v1_.pdf"' in file_res.headers.get("content-disposition")


def test_per_user_limit_50_resumes(client, db_session):
    reg = client.post("/api/auth/register", json={
        "name": "Quota Tester",
        "email": "quota@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201
    user_id = reg.json()["id"]

    for i in range(50):
        r = Resume(
            user_id=user_id,
            filename=f"resume_{i}.pdf",
            stored_path=f"/dummy/path/{i}.pdf",
            mime="application/pdf",
            size=1000,
            extracted_text="dummy text content"
        )
        db_session.add(r)
    db_session.commit()

    pdf_bytes = create_sample_pdf()
    files = {"file": ("overflow.pdf", pdf_bytes, "application/pdf")}
    res = client.post("/api/resumes", files=files)
    assert res.status_code == 409
    assert "limit of 50 resumes reached" in res.json()["detail"].lower()


def test_docx_zip_bomb_rejected(client):
    reg = client.post("/api/auth/register", json={
        "name": "Zip Bomb Tester",
        "email": "zipbomb@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"0" * (55 * 1024 * 1024))
    zip_bomb_bytes = buf.getvalue()

    files = {"file": ("bomb.docx", zip_bomb_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/resumes", files=files)
    assert res.status_code == 422
    assert "exceeds maximum allowed limit of 50mb" in res.json()["detail"].lower() or "suspicious compression ratio" in res.json()["detail"].lower()


def test_streaming_oversize_upload_aborted(client):
    reg = client.post("/api/auth/register", json={
        "name": "Stream Tester",
        "email": "streamtest@example.com",
        "password": "password123"
    })
    assert reg.status_code == 201

    class ChunkStream(io.RawIOBase):
        def __init__(self, total_size, chunk_size=64 * 1024):
            self.total_size = total_size
            self.chunk_size = chunk_size
            self.sent = 0

        def readinto(self, b):
            if self.sent >= self.total_size:
                return 0
            to_read = min(len(b), self.total_size - self.sent)
            b[:to_read] = b"A" * to_read
            self.sent += to_read
            return to_read

    stream_content = ChunkStream(6 * 1024 * 1024)
    files = {"file": ("oversize_stream.pdf", stream_content, "application/pdf")}
    res = client.post("/api/resumes", files=files)
    assert res.status_code == 422
    assert "exceeds maximum allowed limit" in res.json()["detail"].lower()
