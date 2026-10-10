import io
import os
import sys
import time
import uuid
import requests
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document


def generate_pdf_resume() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    lines = [
        "Morgan Lee | morgan.lee@example.com | (555) 234-5678 | github.com/morganlee | linkedin.com/in/morganlee",
        "Professional Summary",
        "Senior Backend Systems Architect with 7+ years designing distributed microservices and reliable data systems.",
        "Technical Skills",
        "Python, FastAPI, Docker, PostgreSQL, Kubernetes, Redis, React, TypeScript, Git, AWS, CI/CD",
        "Professional Experience",
        "Lead Infrastructure Engineer at Apex Data (2021 - Present)",
        "- Architected high-throughput ingestion cluster processing 15M transactions per day.",
        "- Reduced system latency by 42% through connection pooling and Redis caching optimizations.",
        "- Led migration of 12 internal services to Kubernetes clusters with zero downtime.",
        "Backend Software Engineer at CloudPulse (2018 - 2021)",
        "- Developed RESTful microservices in Python and FastAPI with 99.98% uptime SLA.",
        "- Designed relational database schemas in PostgreSQL with automated migration scripts.",
        "Education",
        "Bachelor of Science in Computer Science, University of Washington (2018)",
        "Projects",
        "Distributed Workflow Coordinator (github.com/morganlee/workflow-coord)",
        "Certifications",
        "AWS Certified Solutions Architect Associate"
    ]
    y = 750
    for line in lines:
        if y < 50:
            c.showPage()
            y = 750
        c.drawString(40, y, line[:100])
        y -= 22
    c.save()
    return buf.getvalue()


def generate_docx_resume() -> bytes:
    doc = Document()
    doc.add_heading("Morgan Lee", level=0)
    doc.add_paragraph("morgan.lee@example.com | github.com/morganlee | linkedin.com/in/morganlee")
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph("Senior Backend Systems Architect with 7+ years designing distributed microservices.")
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Python, FastAPI, Docker, PostgreSQL, Kubernetes, Redis, React, TypeScript, Git, AWS, CI/CD")
    doc.add_heading("Professional Experience", level=1)
    doc.add_paragraph("Lead Infrastructure Engineer at Apex Data (2021 - Present)")
    doc.add_paragraph("- Architected high-throughput ingestion cluster processing 15M transactions per day.")
    doc.add_paragraph("- Reduced system latency by 42% through connection pooling and Redis caching optimizations.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science, University of Washington (2018)")
    doc.add_heading("Projects", level=1)
    doc.add_paragraph("Distributed Workflow Coordinator (github.com/morganlee/workflow-coord)")
    doc.add_heading("Certifications", level=1)
    doc.add_paragraph("AWS Certified Solutions Architect Associate")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main():
    base_url = os.environ.get("BASE_URL", "http://127.0.0.1:8000").rstrip("/")
    api_url = f"{base_url}/api"
    print(f"Starting E2E verification against: {base_url}")

    session_a = requests.Session()
    health_resp = session_a.get(f"{api_url}/health")
    assert health_resp.status_code == 200, f"Health check failed: {health_resp.status_code}"
    print("[PASS] Backend health check verified (status: ok)")

    email_a = f"user_a_{uuid.uuid4().hex[:8]}@example.com"
    reg_a = session_a.post(
        f"{api_url}/auth/register",
        json={"name": "User Alpha", "email": email_a, "password": "Password123!"}
    )
    assert reg_a.status_code == 201, f"User A registration failed: {reg_a.text}"
    user_a_id = reg_a.json()["id"]
    print(f"[PASS] Registered User A ({email_a}, ID: {user_a_id})")

    fake_exe_bytes = b"MZ\x90\x00" + b"\x00" * 200
    bad_upload = session_a.post(
        f"{api_url}/resumes",
        files={"file": ("malicious_executable.pdf", fake_exe_bytes, "application/pdf")}
    )
    assert bad_upload.status_code == 422, f"Expected 422 for executable payload, got {bad_upload.status_code}: {bad_upload.text}"
    print("[PASS] Rejected malicious executable disguised as PDF (HTTP 422)")

    pdf_bytes = generate_pdf_resume()
    pdf_upload = session_a.post(
        f"{api_url}/resumes",
        files={"file": ("Morgan_Lee_Resume.pdf", pdf_bytes, "application/pdf")}
    )
    assert pdf_upload.status_code == 201, f"PDF upload failed: {pdf_upload.text}"
    resume_a_pdf_id = pdf_upload.json()["id"]
    print(f"[PASS] Uploaded realistic PDF resume (ID: {resume_a_pdf_id})")

    docx_bytes = generate_docx_resume()
    docx_upload = session_a.post(
        f"{api_url}/resumes",
        files={"file": ("Morgan_Lee_Resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert docx_upload.status_code == 201, f"DOCX upload failed: {docx_upload.text}"
    resume_a_docx_id = docx_upload.json()["id"]
    print(f"[PASS] Uploaded realistic DOCX resume (ID: {resume_a_docx_id})")

    analysis_no_jd_req = session_a.post(
        f"{api_url}/analyses",
        json={"resume_id": resume_a_pdf_id, "job_description": None}
    )
    assert analysis_no_jd_req.status_code == 202, f"Analysis without JD creation failed: {analysis_no_jd_req.text}"
    analysis_1_id = analysis_no_jd_req.json()["id"]
    print(f"[PASS] Queued analysis without JD (ID: {analysis_1_id})")

    analysis_1 = None
    for _ in range(30):
        poll_resp = session_a.get(f"{api_url}/analyses/{analysis_1_id}")
        assert poll_resp.status_code == 200
        data = poll_resp.json()
        if data["status"] == "done":
            analysis_1 = data
            break
        elif data["status"] == "failed":
            raise RuntimeError(f"Analysis 1 failed: {data.get('error')}")
        time.sleep(0.5)

    assert analysis_1 is not None, "Analysis 1 timed out before completing"
    sec1 = analysis_1["sections"]
    expected_sum_1 = round(sum(sec1[s]["score"] * sec1[s]["weight"] for s in ["skills", "experience", "education", "projects", "certifications", "structure"]))
    assert analysis_1["resume_score"] == expected_sum_1, f"Expected {expected_sum_1}, got {analysis_1['resume_score']}"
    assert analysis_1["compatibility_score"] is None, "Expected None compatibility score without JD"
    print(f"[PASS] Analysis without JD complete: score={analysis_1['resume_score']} matches weighted sub-score sum ({expected_sum_1})")

    jd_text = "Required: Python, FastAPI, Docker, PostgreSQL. Preferred: Kubernetes, Redis, React, Rust."
    analysis_with_jd_req = session_a.post(
        f"{api_url}/analyses",
        json={"resume_id": resume_a_docx_id, "job_description": jd_text}
    )
    assert analysis_with_jd_req.status_code == 202, f"Analysis with JD creation failed: {analysis_with_jd_req.text}"
    analysis_2_id = analysis_with_jd_req.json()["id"]
    print(f"[PASS] Queued analysis with JD (ID: {analysis_2_id})")

    analysis_2 = None
    for _ in range(30):
        poll_resp = session_a.get(f"{api_url}/analyses/{analysis_2_id}")
        assert poll_resp.status_code == 200
        data = poll_resp.json()
        if data["status"] == "done":
            analysis_2 = data
            break
        elif data["status"] == "failed":
            raise RuntimeError(f"Analysis 2 failed: {data.get('error')}")
        time.sleep(0.5)

    assert analysis_2 is not None, "Analysis 2 timed out before completing"
    sec2 = analysis_2["sections"]
    expected_sum_2 = round(sum(sec2[s]["score"] * sec2[s]["weight"] for s in ["skills", "experience", "education", "projects", "certifications", "structure"]))
    assert analysis_2["resume_score"] == expected_sum_2, f"Expected {expected_sum_2}, got {analysis_2['resume_score']}"
    assert analysis_2["compatibility_score"] is not None, "Expected non-None compatibility score with JD"
    assert "Python" in analysis_2["matching_skills"]
    assert "Docker" in analysis_2["matching_skills"]
    assert "PostgreSQL" in analysis_2["matching_skills"]
    assert "Rust" in analysis_2["missing_skills"]
    print(f"[PASS] Analysis with JD complete: score={analysis_2['resume_score']} matches weighted sum ({expected_sum_2}), compatibility={analysis_2['compatibility_score']}, matching={analysis_2['matching_skills']}, missing={analysis_2['missing_skills']}")

    history_resp = session_a.get(f"{api_url}/analyses")
    assert history_resp.status_code == 200
    history_ids = [item["id"] for item in history_resp.json()["items"]]
    assert analysis_1_id in history_ids, "Analysis 1 missing from history"
    assert analysis_2_id in history_ids, "Analysis 2 missing from history"
    print(f"[PASS] History listing verified ({len(history_ids)} analyses found including both test runs)")

    file_resp = session_a.get(f"{api_url}/resumes/{resume_a_pdf_id}/file")
    assert file_resp.status_code == 200, f"File preview failed: {file_resp.status_code}"
    assert file_resp.headers.get("x-content-type-options") == "nosniff"
    assert file_resp.headers.get("content-security-policy") == "sandbox"
    assert "inline; filename=" in file_resp.headers.get("content-disposition", "")
    assert len(file_resp.content) == len(pdf_bytes)
    print("[PASS] File preview and download verified with nosniff, sandbox, and inline filename headers")

    session_b = requests.Session()
    email_b = f"user_b_{uuid.uuid4().hex[:8]}@example.com"
    reg_b = session_b.post(
        f"{api_url}/auth/register",
        json={"name": "User Beta", "email": email_b, "password": "Password123!"}
    )
    assert reg_b.status_code == 201
    print(f"[PASS] Registered User B ({email_b})")

    b_get_resume = session_b.get(f"{api_url}/resumes/{resume_a_pdf_id}")
    assert b_get_resume.status_code == 404, f"Expected 404, got {b_get_resume.status_code}"

    b_get_file = session_b.get(f"{api_url}/resumes/{resume_a_pdf_id}/file")
    assert b_get_file.status_code == 404, f"Expected 404, got {b_get_file.status_code}"

    b_get_analysis1 = session_b.get(f"{api_url}/analyses/{analysis_1_id}")
    assert b_get_analysis1.status_code == 404, f"Expected 404, got {b_get_analysis1.status_code}"

    b_get_analysis2 = session_b.get(f"{api_url}/analyses/{analysis_2_id}")
    assert b_get_analysis2.status_code == 404, f"Expected 404, got {b_get_analysis2.status_code}"
    print("[PASS] Multi-tenant isolation verified: User B gets 404 for all User A resumes, files, and analyses")

    print("\nALL E2E VERIFICATIONS PASSED SUCCESSFULLY.")


if __name__ == "__main__":
    main()
